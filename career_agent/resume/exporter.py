"""Generate polished Word and PDF downloads from structured resume data."""

from __future__ import annotations

import base64
import re
from collections.abc import Iterable
from dataclasses import dataclass
from html import escape
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage
from PIL import UnidentifiedImageError
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer

from career_agent.resume.editor import normalize_section_order
from career_agent.resume.models import Resume, TailoredResume
from career_agent.resume.studio import ResumeDesignTheme


@dataclass(frozen=True)
class ResumeTheme:
    font_name: str
    east_asia_font: str
    accent: str
    title_size: float
    body_size: float
    heading_size: float
    title_centered: bool
    compact: bool = False
    layout: str = "editorial"
    visual_style: str = "auto"


THEMES = {
    "modern": ResumeTheme(
        "Arial", "Microsoft YaHei", "176B72", 26, 10.5, 13, False, False, "banner"
    ),
    "classic": ResumeTheme("Times New Roman", "SimSun", "202020", 25, 10.5, 13, True),
    "minimal": ResumeTheme("Arial", "DengXian", "56606A", 24, 10, 11.5, False, True, "sidebar"),
}


def section_overrides_with_suggestions(
    tailored: TailoredResume, selected_indices: list[int] | None = None
) -> dict[str, list[str]]:
    """Include unverified draft lines only when explicitly selected for this download."""

    selected = set(selected_indices or [])
    if any(index < 0 or index >= len(tailored.suggestions) for index in selected):
        raise ValueError("所选补写句子已过期，请重新选择")
    sections = {
        name: [bullet.text for bullet in bullets] for name, bullets in tailored.sections.items()
    }
    for index in sorted(selected):
        suggestion = tailored.suggestions[index]
        sections.setdefault(suggestion.section, []).append(suggestion.text)
    return sections


def render_resume_docx(
    resume: Resume,
    *,
    hometown: str | None = None,
    photo_asset: str | None = None,
    template_id: str = "modern",
    section_order: list[str] | None = None,
    custom_template: bytes | None = None,
    custom_theme: dict | None = None,
    section_overrides: dict[str, list[str]] | None = None,
) -> bytes:
    theme = _theme(template_id, custom_template, custom_theme)
    document = _document_from_template(custom_template)
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    if not custom_template:
        margin = 0.62 if theme.compact else 0.72
        section.top_margin = Inches(margin)
        section.bottom_margin = Inches(margin)
        section.left_margin = Inches(0.78)
        section.right_margin = Inches(0.78)

    _configure_word_styles(document, theme)
    photo = _photo_bytes(photo_asset)
    content = document
    visual_style = _visual_style(theme)
    if theme.layout == "sidebar":
        table = document.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width = Inches(1.65)
        table.columns[1].width = Inches(4.95)
        _word_cell_color(table.cell(0, 0), "102E51" if visual_style == "blueprint" else "E8F0EF")
        _word_identity(
            table.cell(0, 0), resume, hometown, photo, theme, inverse=visual_style == "blueprint"
        )
        content = table.cell(0, 1)
    elif theme.layout == "banner":
        table = document.add_table(rows=1, cols=2 if photo else 1)
        _word_cell_color(table.cell(0, 0), "0B2444" if visual_style == "aurora" else theme.accent)
        _word_identity(table.cell(0, 0), resume, hometown, None, theme, inverse=True)
        if photo:
            _word_cell_color(
                table.cell(0, 1), "0B2444" if visual_style == "aurora" else theme.accent
            )
            table.cell(0, 1).add_paragraph().add_run().add_picture(photo, width=Inches(0.72))
        document.add_paragraph().paragraph_format.space_after = Pt(2)
    else:
        _word_identity(document, resume, hometown, photo, theme)

    for heading, entries in _resume_sections(resume, section_order, section_overrides):
        heading_paragraph = content.add_paragraph(style="Heading 1")
        heading_paragraph.add_run(_safe_text(heading))
        _word_heading_rule(heading_paragraph, theme.accent)
        for entry_title, description in entries:
            paragraph = content.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.add_run(_safe_text(entry_title)).bold = True
            if description:
                detail = content.add_paragraph(_safe_text(description))
                detail.paragraph_format.left_indent = Inches(0.08)
                detail.paragraph_format.space_after = Pt(7 if not theme.compact else 4)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_resume_pdf(
    resume: Resume,
    *,
    hometown: str | None = None,
    photo_asset: str | None = None,
    template_id: str = "modern",
    section_order: list[str] | None = None,
    custom_template: bytes | None = None,
    custom_theme: dict | None = None,
    section_overrides: dict[str, list[str]] | None = None,
) -> bytes:
    theme = _theme(template_id, custom_template, custom_theme)
    visual_style = _visual_style(theme)
    regular_font, bold_font = _register_pdf_fonts(theme)
    output = BytesIO()
    margin = 16 * mm if theme.compact else 19 * mm
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=70 * mm if theme.layout == "sidebar" else margin,
        rightMargin=margin,
        topMargin=54 * mm if theme.layout == "banner" else 16 * mm,
        bottomMargin=16 * mm,
        title=f"{resume.personal_info.name or '个人'}简历",
        author="",
    )
    base = getSampleStyleSheet()
    accent = f"#{theme.accent}"
    title_alignment = TA_CENTER if theme.title_centered else TA_LEFT
    title_style = ParagraphStyle(
        "ResumeTitle",
        parent=base["Title"],
        fontName=bold_font,
        fontSize=theme.title_size,
        leading=theme.title_size + 7,
        textColor="#17212B",
        alignment=title_alignment,
        spaceAfter=5,
    )
    contact_style = ParagraphStyle(
        "Contact",
        parent=base["Normal"],
        fontName=regular_font,
        fontSize=9.2,
        leading=14,
        alignment=title_alignment,
        textColor="#56606A",
        spaceAfter=12,
        wordWrap="CJK",
    )
    heading_style = ParagraphStyle(
        "ResumeHeading",
        parent=base["Heading2"],
        fontName=bold_font,
        fontSize=theme.heading_size,
        leading=18,
        textColor=accent,
        spaceBefore=10 if theme.compact else 14,
        spaceAfter=6,
    )
    entry_style = ParagraphStyle(
        "ResumeEntry",
        parent=base["BodyText"],
        fontName=bold_font,
        fontSize=theme.body_size,
        leading=16,
        textColor="#1F2933",
        spaceAfter=2,
        wordWrap="CJK",
    )
    body_style = ParagraphStyle(
        "ResumeBody",
        parent=base["BodyText"],
        fontName=regular_font,
        fontSize=theme.body_size,
        leading=16.5,
        textColor="#303B44",
        leftIndent=4,
        spaceAfter=6 if theme.compact else 9,
        wordWrap="CJK",
    )
    story: list[object] = []
    if theme.layout == "editorial":
        story.extend(
            (
                Paragraph(escape(_safe_text(resume.personal_info.name or "个人简历")), title_style),
                Paragraph(escape(_safe_text(_contact_text(resume, hometown))), contact_style),
            )
        )
        if resume.personal_info.links:
            story.append(
                Paragraph(
                    escape(_safe_text(_profile_links_text(resume))),
                    ParagraphStyle("ProfileLinks", parent=contact_style, fontSize=8.4),
                )
            )
    photo = _photo_bytes(photo_asset)
    if photo and theme.layout == "editorial":
        image = Image(photo, width=18 * mm, height=24 * mm)
        image.hAlign = "RIGHT"
        story.extend((image, Spacer(1, 3 * mm)))
    for heading, entries in _resume_sections(resume, section_order, section_overrides):
        story.append(Paragraph(escape(_safe_text(heading)), heading_style))
        for entry_title, description in entries:
            story.append(Paragraph(escape(_safe_text(entry_title)), entry_style))
            if description:
                story.append(Paragraph(escape(_safe_text(description)), body_style))

    def page_design(canvas, doc):
        canvas.saveState()
        if theme.layout == "sidebar":
            dark = visual_style == "blueprint"
            canvas.setFillColor(HexColor("#102E51" if dark else "#E8F0EF"))
            canvas.rect(0, 0, 62 * mm, A4[1], fill=1, stroke=0)
            if dark:
                canvas.setStrokeColor(HexColor("#1D4F69"))
                canvas.setLineWidth(0.25)
                for position in range(0, 62, 7):
                    canvas.line(position * mm, 0, position * mm, A4[1])
                for position in range(0, 300, 7):
                    canvas.line(0, position * mm, 62 * mm, position * mm)
                canvas.setStrokeColor(HexColor("#319B9B"))
                canvas.circle(-8 * mm, 54 * mm, 44 * mm, fill=0, stroke=1)
                canvas.setFillColor(HexColor(accent))
                canvas.rect(61 * mm, 0, 1 * mm, A4[1], fill=1, stroke=0)
            side_title = ParagraphStyle(
                "SidebarTitle", parent=title_style, textColor="#FFFFFF" if dark else "#17212B"
            )
            side_contact = ParagraphStyle(
                "SidebarContact", parent=contact_style, textColor="#D5ECF2" if dark else "#56606A"
            )
            _pdf_identity_on_canvas(
                canvas,
                resume,
                hometown,
                photo,
                side_title,
                side_contact,
                x=9 * mm,
                width=48 * mm,
                top=A4[1] - 25 * mm,
            )
        elif theme.layout == "banner" and doc.page == 1:
            if visual_style == "aurora":
                colors = ("#0B1E3C", "#103C5C", "#155C70", "#178F8B")
                band = A4[0] / len(colors)
                for index, color in enumerate(colors):
                    canvas.setFillColor(HexColor(color))
                    canvas.rect(index * band, A4[1] - 44 * mm, band + 1, 44 * mm, fill=1, stroke=0)
                canvas.setStrokeColor(HexColor("#80D9D7"))
                canvas.circle(A4[0] - 18 * mm, A4[1] - 13 * mm, 28 * mm, fill=0, stroke=1)
                canvas.setFillColor(HexColor("#8CEBDB"))
                canvas.rect(0, A4[1] - 44 * mm, A4[0], 1.2 * mm, fill=1, stroke=0)
            else:
                canvas.setFillColor(HexColor(accent))
                canvas.rect(0, A4[1] - 44 * mm, A4[0], 44 * mm, fill=1, stroke=0)
            banner_title = ParagraphStyle("BannerTitle", parent=title_style, textColor="#FFFFFF")
            banner_contact = ParagraphStyle(
                "BannerContact", parent=contact_style, textColor="#FFFFFF"
            )
            _pdf_identity_on_canvas(
                canvas,
                resume,
                hometown,
                photo,
                banner_title,
                banner_contact,
                x=19 * mm,
                width=145 * mm,
                top=A4[1] - 9 * mm,
                photo_x=A4[0] - 40 * mm,
                photo_y=A4[1] - 34 * mm,
            )
        elif theme.layout == "editorial" and visual_style == "paper":
            canvas.setFillColor(HexColor("#FCFAF5"))
            canvas.rect(0, 0, A4[0], A4[1], fill=1, stroke=0)
            canvas.setStrokeColor(HexColor("#E3D8C8"))
            canvas.rect(10 * mm, 10 * mm, A4[0] - 20 * mm, A4[1] - 20 * mm, fill=0, stroke=1)
        canvas.restoreState()

    document.build(story, onFirstPage=page_design, onLaterPages=page_design)
    return output.getvalue()


def inspect_docx_template(content: bytes) -> dict[str, str | float | bool]:
    """Extract a small, safe visual theme from an uploaded Word template."""

    document = Document(BytesIO(content))
    normal = document.styles["Normal"].font
    title = document.styles["Title"].font
    heading = document.styles["Heading 1"].font
    color = heading.color.rgb
    return {
        "font_family": normal.name or "Microsoft YaHei",
        "title_size": title.size.pt if title.size else 26,
        "heading_size": heading.size.pt if heading.size else 13,
        "accent": f"#{color}" if color else "#176B72",
        "title_centered": False,
    }


def _document_from_template(content: bytes | None) -> Document:
    if not content:
        return Document()
    document = Document(BytesIO(content))
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    return document


def _theme(
    template_id: str, custom_template: bytes | None, custom_theme: dict | None = None
) -> ResumeTheme:
    if template_id == "designed" and custom_theme:
        values = ResumeDesignTheme.model_validate(custom_theme)
        return ResumeTheme(
            values.font_family,
            values.font_family,
            values.accent.lstrip("#"),
            values.title_size,
            values.body_size,
            values.heading_size,
            values.title_centered,
            values.compact,
            values.layout,
            values.visual_style,
        )
    if template_id == "custom" and custom_template:
        values = inspect_docx_template(custom_template)
        return ResumeTheme(
            str(values["font_family"]),
            str(values["font_family"]),
            str(values["accent"]).lstrip("#"),
            float(values["title_size"]),
            10.5,
            float(values["heading_size"]),
            bool(values["title_centered"]),
        )
    return THEMES.get(template_id, THEMES["modern"])


def _configure_word_styles(document: Document, theme: ResumeTheme) -> None:
    for style_name, size, bold, color in (
        ("Normal", theme.body_size, False, "303B44"),
        ("Title", theme.title_size, True, "17212B"),
        ("Heading 1", theme.heading_size, True, theme.accent),
    ):
        style = document.styles[style_name]
        style.font.name = theme.font_name
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), theme.east_asia_font)
    document.styles["Normal"].paragraph_format.line_spacing = 1.25
    document.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(5)


def _visual_style(theme: ResumeTheme) -> str:
    if theme.visual_style != "auto":
        return theme.visual_style
    return {"sidebar": "blueprint", "banner": "aurora"}.get(theme.layout, "paper")


def _word_identity(
    container: object,
    resume: Resume,
    hometown: str | None,
    photo: bytes | None,
    theme: ResumeTheme,
    *,
    inverse: bool = False,
) -> None:
    title = container.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER if theme.title_centered else WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run(_safe_text(resume.personal_info.name or "个人简历"))
    if inverse:
        title_run.font.color.rgb = RGBColor(255, 255, 255)
    contact = container.add_paragraph(_safe_text(_contact_text(resume, hometown)))
    contact.alignment = title.alignment
    contact.paragraph_format.space_after = Pt(10)
    if inverse:
        for run in contact.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
    if resume.personal_info.links:
        links = container.add_paragraph(_safe_text(_profile_links_text(resume, separator="\n")))
        links.paragraph_format.space_after = Pt(7)
        if inverse:
            for run in links.runs:
                run.font.color.rgb = RGBColor(222, 244, 247)
    if photo:
        paragraph = container.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.add_run().add_picture(photo, width=Inches(0.72))


def _word_cell_color(cell: object, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _pdf_identity_on_canvas(
    canvas: object,
    resume: Resume,
    hometown: str | None,
    photo: bytes | None,
    title_style: ParagraphStyle,
    contact_style: ParagraphStyle,
    *,
    x: float,
    width: float,
    top: float,
    photo_x: float | None = None,
    photo_y: float | None = None,
) -> None:
    name = Paragraph(escape(_safe_text(resume.personal_info.name or "个人简历")), title_style)
    _, title_height = name.wrap(width, 50 * mm)
    name.drawOn(canvas, x, top - title_height)
    contact = Paragraph(escape(_safe_text(_contact_text(resume, hometown))), contact_style)
    _, contact_height = contact.wrap(width, 90 * mm)
    contact.drawOn(canvas, x, top - title_height - 4 * mm - contact_height)
    links_height = 0
    if resume.personal_info.links:
        links_style = ParagraphStyle(
            "CanvasLinks", parent=contact_style, fontSize=7.8, leading=10, spaceAfter=0
        )
        link_text = escape(_safe_text(_profile_links_text(resume, separator="\n")))
        links = Paragraph(link_text.replace("\n", "<br/>"), links_style)
        _, links_height = links.wrap(width, 90 * mm)
        links.drawOn(canvas, x, top - title_height - 5 * mm - contact_height - links_height)
    if photo:
        from reportlab.lib.utils import ImageReader

        canvas.drawImage(
            ImageReader(photo),
            photo_x if photo_x is not None else x,
            photo_y
            if photo_y is not None
            else (top - title_height - contact_height - links_height - 35 * mm),
            18 * mm,
            24 * mm,
            preserveAspectRatio=True,
            mask="auto",
        )


def _word_heading_rule(paragraph: object, color: str) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    paragraph_properties.append(borders)


def _contact_text(resume: Resume, hometown: str | None) -> str:
    info = resume.personal_info
    values = [value for value in (info.email, info.phone) if value]
    if info.location:
        values.append(f"所在地：{info.location}")
    if hometown or info.hometown:
        values.append(f"籍贯：{hometown or info.hometown}")
    return " · ".join(values)


def _profile_links_text(resume: Resume, *, separator: str = " · ") -> str:
    return separator.join(
        f"{'GitHub' if 'github.com' in link.casefold() else '个人网站'}：{link}"
        for link in resume.personal_info.links
    )


def _resume_sections(
    resume: Resume,
    section_order: list[str] | None,
    section_overrides: dict[str, list[str]] | None = None,
) -> Iterable[tuple[str, list[tuple[str, str]]]]:
    sections: dict[str, tuple[str, list[tuple[str, str]]]] = {}
    if resume.summary:
        sections["summary"] = ("个人总结", [(resume.summary, "")])
    if resume.educations:
        sections["education"] = (
            "教育经历",
            [
                (
                    " · ".join(
                        value
                        for value in (
                            item.school,
                            item.major,
                            item.degree,
                            _dates(item.start_date, item.end_date),
                        )
                        if value
                    ),
                    item.description or "",
                )
                for item in resume.educations
            ],
        )
    for key, heading, items, name_field, role_field in (
        ("work", "工作经历", resume.work_experiences, "company", "title"),
        ("internship", "实习经历", resume.internships, "company", "title"),
        ("experience", "经历", resume.project_experiences, "name", "role"),
        ("campus", "校园经历", resume.campus_experiences, "organization", "role"),
    ):
        if items:
            entries = []
            for item in items:
                title = " · ".join(
                    value
                    for value in (
                        getattr(item, name_field) or heading,
                        getattr(item, role_field) or "",
                        _dates(item.start_date, item.end_date),
                    )
                    if value
                )
                description = item.description or ""
                if key == "experience" and item.technologies:
                    technologies = "相关技术：" + "、".join(item.technologies)
                    description = "\n".join(part for part in (description, technologies) if part)
                entries.append((title, description))
            sections[key] = (heading, entries)
    if resume.skills:
        sections["skills"] = (
            "专业技能",
            [("、".join(item.name for item in resume.skills), "")],
        )
    aliases = {"project": "experience", "skill": "skills"}
    headings = {
        "summary": "个人总结",
        "education": "教育经历",
        "work": "工作经历",
        "internship": "实习经历",
        "experience": "经历",
        "campus": "校园经历",
        "skills": "专业技能",
    }
    for original_key, bullets in (section_overrides or {}).items():
        key = aliases.get(original_key, original_key)
        values = [value.strip() for value in bullets if value.strip()]
        if not values:
            continue
        heading, existing = sections.get(key, (headings.get(key, key), []))
        if key == "skills":
            entries = [("、".join(values), "")]
        elif len(existing) == len(values):
            entries = [(title, value) for (title, _), value in zip(existing, values, strict=True)]
        elif existing:
            entries = [(existing[0][0], "\n".join(f"• {value}" for value in values))]
            entries.extend(existing[1:])
        else:
            entries = [(value, "") for value in values]
        sections[key] = (heading, entries)
    for key in normalize_section_order(section_order or []):
        if key in sections:
            heading, entries = sections[key]
            yield resume.section_titles.get(key, heading), entries


def _dates(start: str | None, end: str | None) -> str:
    return " 至 ".join(value for value in (start, end) if value)


def _photo_bytes(photo_asset: str | None) -> BytesIO | None:
    if not photo_asset or not photo_asset.startswith("data:image/"):
        return None
    try:
        header, encoded = photo_asset.split(",", 1)
        if ";base64" not in header:
            return None
        source = BytesIO(base64.b64decode(encoded, validate=True))
        with PILImage.open(source) as image:
            image.load()
            converted = image.convert("RGB")
            converted.thumbnail((1200, 1600))
            output = BytesIO()
            converted.save(output, format="PNG", optimize=True)
            output.seek(0)
            return output
    except (ValueError, TypeError, OSError, UnidentifiedImageError):
        return None


def _safe_text(value: object) -> str:
    """Remove characters rejected by OOXML and ReportLab while preserving user text."""

    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", str(value))


def _register_pdf_fonts(theme: ResumeTheme) -> tuple[str, str]:
    if theme.east_asia_font in {"SimSun", "Georgia"}:
        family = "CareerResume-Serif"
        regular_path = Path("C:/Windows/Fonts/simsun.ttc")
        bold_path = Path("C:/Windows/Fonts/simhei.ttf")
    elif theme.east_asia_font == "Microsoft YaHei":
        family = "CareerResume-YaHei"
        regular_path = Path("C:/Windows/Fonts/msyh.ttc")
        bold_path = Path("C:/Windows/Fonts/msyhbd.ttc")
    else:
        family = "CareerResume-Deng"
        regular_path = Path("C:/Windows/Fonts/Deng.ttf")
        bold_path = Path("C:/Windows/Fonts/Dengb.ttf")
    if regular_path.exists() and bold_path.exists():
        if family not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(family, regular_path))
            pdfmetrics.registerFont(TTFont(f"{family}-Bold", bold_path))
            pdfmetrics.registerFontFamily(
                family,
                normal=family,
                bold=f"{family}-Bold",
                italic=family,
                boldItalic=f"{family}-Bold",
            )
        return family, f"{family}-Bold"
    if "STSong-Light" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light", "STSong-Light"
