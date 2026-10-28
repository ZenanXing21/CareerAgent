import base64
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from career_agent import CandidateIntakeDraft, EducationDraft
from career_agent.job_platform.liepin_mcp import LiepinMCPAdapter
from career_agent.llm.openai_compatible import OpenAICompatibleProvider
from career_agent.resume.models import CampusExperience, ExperienceType, Resume, ResumeAgentEdit
from career_agent.web.app import (
    IntakeAttachment,
    IntakeConversationSession,
    JobApplyRequest,
    ResumeAgentRequest,
    ResumeEditRequest,
    _decode_attachment,
    _next_intake_message,
    create_app,
)
from career_agent.web.workspace import CandidateRecord, WorkspaceRecords, WorkspaceStore


def test_web_app_serves_workspace_and_reports_safe_configuration(monkeypatch) -> None:
    monkeypatch.setenv("CAREER_AGENT_LLM_API_KEY", "test-key")
    monkeypatch.setenv("CAREER_AGENT_LLM_MODEL", "test-model")
    monkeypatch.setenv("LIEPIN_MCP_USER_TOKEN", "test-token")
    client = TestClient(create_app())

    page = client.get("/")
    status = client.get("/api/status")

    assert page.status_code == 200
    assert "Career Agent 工作台" in page.text
    assert "resume-agent-form" in page.text
    assert "copy-job-jd" in page.text
    assert "analysis-requirements" in page.text
    assert "intake-chat-form" in page.text
    assert "intake-form" not in page.text
    assert status.json() == {"ready": True, "model": "test-model", "liepin": True}
    assert "test-key" not in status.text
    assert "test-token" not in status.text


def test_resume_edit_payload_accepts_json_enum_values() -> None:
    payload = {
        "resume_id": "resume-web",
        "version": "v1",
        "raw_text": "张明\n工作：测试公司；工程师",
        "personal_info": {"name": "张明", "links": []},
        "educations": [],
        "work_experiences": [
            {
                "id": "work_001",
                "source_text": "工作：测试公司；工程师",
                "experience_type": "work",
                "company": "测试公司",
                "title": "工程师",
            }
        ],
        "internships": [],
        "project_experiences": [],
        "campus_experiences": [],
        "skills": [],
    }

    request = ResumeEditRequest(candidate_id="candidate", resume=payload)
    resume = Resume.model_validate(request.resume, strict=False)

    assert resume.work_experiences[0].experience_type is ExperienceType.WORK


def test_job_apply_requires_explicit_attachment_confirmation() -> None:
    request = JobApplyRequest(candidate_id="candidate", job_ids=["liepin:123"])

    assert request.attachment_confirmed is False


def test_resume_agent_request_defaults_to_polished_resume() -> None:
    request = ResumeAgentRequest(candidate_id="candidate", message="补充 SQL 技能")

    assert request.variant == "polished"


def test_resume_agent_contact_changes_are_visible_and_reported(tmp_path, monkeypatch) -> None:
    source = CandidateIntakeDraft(
        candidate_id="profile-edit", name="张明", email="z@example.com",
        phone="13800000000", education=EducationDraft(
            school="测试大学", degree="本科", major="计算机"
        ), skills=["Python"],
    ).build_resume().model_copy(update={"campus_experiences": [CampusExperience(
        id="campus_001", source_text="参与校园活动", organization="校园活动",
        description="参与校园活动",
    )]})
    state_file = tmp_path / "workspace.json"
    WorkspaceStore(state_file).write(WorkspaceRecords(candidates={
        "profile-edit": CandidateRecord(raw_resume=source, display_resume=source)
    }).model_dump_json())

    class EditingProvider:
        async def edit_resume(self, *, resume, instruction, conversation):
            return ResumeAgentEdit(assistant_message="已更新。", resume=resume)

    monkeypatch.setattr(OpenAICompatibleProvider, "from_env", lambda: EditingProvider())
    monkeypatch.setattr(LiepinMCPAdapter, "from_env", lambda: object())
    client = TestClient(create_app(state_file=state_file))
    response = client.post("/api/resume/agent", json={
        "candidate_id": "profile-edit",
        "message": "个人网站 https://example.com，GitHub https://github.com/example，"
                   "所在地是上海，籍贯是江苏苏州。",
    })

    assert response.status_code == 200
    result = response.json()
    info = result["resume"]["personal_info"]
    assert info["links"] == ["https://example.com", "https://github.com/example"]
    assert info["location"] == "上海"
    assert info["hometown"] == "江苏苏州"
    assert set(result["changed_sections"]) == {"个人网站与链接", "所在地", "籍贯"}
    saved = WorkspaceStore(state_file).load().candidates["profile-edit"]
    assert saved.polished_hometown == "江苏苏州"

    correction = client.post("/api/resume/agent", json={
        "candidate_id": "profile-edit", "message": "准确来说这是社会实践经历",
    })
    assert correction.status_code == 200
    assert correction.json()["resume"]["section_titles"]["campus"] == "社会实践经历"
    assert "栏目标题" in correction.json()["changed_sections"]

    manually_edited = correction.json()["resume"]
    manually_edited["summary"] = "认真完成交付。"
    manually_edited["section_titles"]["summary"] = "职业概述"
    saved_edit = client.post("/api/resume/edit", json={
        "candidate_id": "profile-edit", "resume": manually_edited,
    })
    assert saved_edit.status_code == 200
    assert saved_edit.json()["resume"]["section_titles"]["summary"] == "职业概述"
    restored = client.get("/api/resume/current?candidate_id=profile-edit")
    assert restored.json()["display_resume"]["section_titles"]["summary"] == "职业概述"

    downloaded = client.get("/api/resume/download/polished/docx?candidate_id=profile-edit")
    assert downloaded.status_code == 200
    doc = Document(BytesIO(downloaded.content))
    assert any(paragraph.text == "社会实践经历" for paragraph in doc.paragraphs)
    assert any(paragraph.text == "职业概述" for paragraph in doc.paragraphs)


def test_conversational_intake_reports_required_then_optional_information() -> None:
    session = IntakeConversationSession(draft=CandidateIntakeDraft(candidate_id="candidate"))
    message, missing, ready = _next_intake_message(session)
    assert "姓名" in message
    assert missing
    assert ready is False

    complete = CandidateIntakeDraft(
        candidate_id="candidate",
        name="张明",
        email="z@example.com",
        phone="13800000000",
        education=EducationDraft(
            school="测试大学",
            degree="本科",
            major="计算机",
            start_date="2022-09",
            end_date="2026-06",
        ),
        skills=["Python"],
    )
    session.draft = complete
    message, missing, ready = _next_intake_message(session)
    assert missing == []
    assert "可选信息（1/7）" in message
    assert ready is False

    session.phase = "final"
    message, missing, ready = _next_intake_message(session)
    assert missing == []
    assert "还有什么需要补充" in message
    assert "确认生成" in message
    assert ready is True


def test_required_information_precedes_experience_clarification() -> None:
    session = IntakeConversationSession(
        draft=CandidateIntakeDraft(candidate_id="candidate"),
        pending_clarification="这段经历是正式工作还是实习？",
    )

    message, missing, ready = _next_intake_message(session)

    assert "姓名" in message
    assert missing
    assert ready is False


def test_ambiguous_work_clarification_precedes_optional_questions() -> None:
    session = IntakeConversationSession(
        draft=CandidateIntakeDraft(
            candidate_id="candidate",
            name="张明",
            email="z@example.com",
            phone="13800000000",
            education=EducationDraft(
                school="测试大学",
                degree="本科",
                major="计算机",
                start_date="2022-09",
                end_date="2026-06",
            ),
            skills=["Python"],
        ),
        pending_clarification="这段经历是正式工作还是实习？",
    )

    message, missing, ready = _next_intake_message(session)

    assert message == "这段经历是正式工作还是实习？"
    assert missing == []
    assert ready is False


def test_conversational_intake_extracts_word_attachment_text() -> None:
    document = Document()
    document.add_paragraph("张明 Python 开发")
    stream = BytesIO()
    document.save(stream)
    encoded = base64.b64encode(stream.getvalue()).decode("ascii")

    text, image = _decode_attachment(
        IntakeAttachment(
            filename="resume.docx",
            data_url=f"data:application/vnd.openxmlformats-officedocument.wordprocessingml.document;base64,{encoded}",
        )
    )

    assert text == "张明 Python 开发"
    assert image is None
