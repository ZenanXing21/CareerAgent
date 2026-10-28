import base64
from io import BytesIO

import pytest
from docx import Document
from PIL import Image

from career_agent import CandidateIntakeDraft, EducationDraft, ExperienceDraft
from career_agent.resume.exporter import (
    inspect_docx_template,
    render_resume_docx,
    render_resume_pdf,
)
from career_agent.resume.models import ProjectExperience


def _resume():  # type: ignore[no-untyped-def]
    return CandidateIntakeDraft(
        candidate_id="download-test",
        name="张明",
        email="zhang@example.com",
        phone="13800000000",
        current_city="上海",
        hometown="苏州",
        education=EducationDraft(school="测试大学", degree="本科", major="计算机"),
        skills=["Python", "SQL"],
        experiences=[ExperienceDraft(content="使用 Python 完成数据清洗")],
    ).build_resume()


def test_word_export_is_a_real_readable_docx() -> None:
    content = render_resume_docx(_resume(), hometown="苏州")
    document = Document(BytesIO(content))

    assert content.startswith(b"PK")
    paragraphs = list(document.paragraphs)
    paragraphs.extend(
        paragraph
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )
    assert any("张明" in paragraph.text for paragraph in paragraphs)


def test_pdf_export_is_a_real_pdf() -> None:
    content = render_resume_pdf(_resume(), hometown="苏州")

    assert content.startswith(b"%PDF-")
    assert len(content) > 1000


def test_preset_layouts_change_word_structure_and_keep_pdf_downloadable() -> None:
    resume = _resume()
    for template_id in ("modern", "classic", "minimal"):
        document = Document(BytesIO(render_resume_docx(resume, template_id=template_id)))
        if template_id == "classic":
            assert any("张明" in item.text for item in document.paragraphs)
        else:
            assert any(
                "张明" in paragraph.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
                for paragraph in cell.paragraphs
            )
        if template_id == "minimal":
            assert any(
                "教育经历" in paragraph.text
                for paragraph in document.tables[0].cell(0, 1).paragraphs
            )
            assert 'w:fill="102E51"' in document.tables[0].cell(0, 0)._tc.xml
        assert render_resume_pdf(resume, template_id=template_id).startswith(b"%PDF-")


def test_generated_preset_can_choose_a_plain_or_decorated_background() -> None:
    resume = _resume()
    for visual_style, expected_fill in (("plain", "E8F0EF"), ("blueprint", "102E51")):
        content = render_resume_docx(
            resume,
            template_id="designed",
            custom_theme={"layout": "sidebar", "visual_style": visual_style},
        )
        document = Document(BytesIO(content))
        assert f'w:fill="{expected_fill}"' in document.tables[0].cell(0, 0)._tc.xml


def test_profile_links_are_present_in_downloaded_resume() -> None:
    resume = _resume()
    info = resume.personal_info.model_copy(
        update={"links": ["https://portfolio.example", "https://github.com/example"]}
    )
    resume = resume.model_copy(update={"personal_info": info})
    document = Document(BytesIO(render_resume_docx(resume, hometown="江苏苏州")))
    text = "\n".join(
        paragraph.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )
    assert "https://portfolio.example" in text
    assert "https://github.com/example" in text
    assert "所在地：上海" in text
    assert "籍贯：江苏苏州" in text


def test_export_respects_section_order_and_custom_word_template() -> None:
    template = Document()
    template.styles["Normal"].font.name = "Arial"
    template_stream = BytesIO()
    template.save(template_stream)

    content = render_resume_docx(
        _resume(),
        section_order=["skills", "education", "experience"],
        template_id="custom",
        custom_template=template_stream.getvalue(),
    )
    document = Document(BytesIO(content))
    text = [paragraph.text for paragraph in document.paragraphs]

    assert text.index("专业技能") < text.index("教育经历") < text.index("经历")
    assert inspect_docx_template(template_stream.getvalue())["font_family"] == "Arial"


def test_word_and_pdf_export_use_edited_section_titles() -> None:
    resume = _resume().model_copy(
        update={
            "summary": "善于沟通与交付。",
            "section_titles": {"summary": "职业概述", "skills": "技术能力"},
        }
    )
    word = render_resume_docx(resume)
    text = "\n".join(p.text for p in Document(BytesIO(word)).paragraphs)

    assert "职业概述" in text
    assert "技术能力" in text
    assert "个人总结" not in text
    assert render_resume_pdf(resume).startswith(b"%PDF-")


def test_word_export_keeps_project_technologies_visible_in_editor() -> None:
    resume = _resume().model_copy(
        update={
            "project_experiences": [
                ProjectExperience(
                    id="project_001",
                    source_text="Agent 项目",
                    name="Agent 项目",
                    description="构建业务助手",
                    technologies=["Python", "FastAPI"],
                )
            ]
        }
    )
    word = render_resume_docx(resume)
    document = Document(BytesIO(word))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "相关技术：Python、FastAPI" in text


def test_word_export_applies_tailored_section_content() -> None:
    content = render_resume_docx(
        _resume(),
        section_overrides={
            "project": ["针对 Agent 岗位突出 Python 数据处理能力"],
            "skills": ["Python 与 SQL 工程实践"],
        },
    )
    text = "\n".join(paragraph.text for paragraph in Document(BytesIO(content)).paragraphs)

    assert "针对 Agent 岗位突出 Python 数据处理能力" in text
    assert "Python 与 SQL 工程实践" in text


def test_unverified_suggestions_require_explicit_download_selection() -> None:
    from career_agent.resume.exporter import section_overrides_with_suggestions
    from career_agent.resume.models import DraftSuggestion, TailoredBullet, TailoredResume

    tailored = TailoredResume(
        source_resume_id="r",
        source_resume_version="v",
        target_job_id="j",
        sections={
            "project": [TailoredBullet(text="已证实的项目经验", source_fact_ids=["fact_001"])]
        },
        suggestions=[DraftSuggestion(section="project", text="待人工核实的补写句子")],
    )
    verified = section_overrides_with_suggestions(tailored)
    reviewed = section_overrides_with_suggestions(tailored, [0])

    assert verified == {"project": ["已证实的项目经验"]}
    assert reviewed == {"project": ["已证实的项目经验", "待人工核实的补写句子"]}
    assert "待人工核实的补写句子" not in "\n".join(
        paragraph.text
        for paragraph in Document(
            BytesIO(render_resume_docx(_resume(), section_overrides=verified))
        ).paragraphs
    )
    assert "待人工核实的补写句子" in "\n".join(
        paragraph.text
        for paragraph in Document(
            BytesIO(render_resume_docx(_resume(), section_overrides=reviewed))
        ).paragraphs
    )
    with pytest.raises(ValueError, match="已过期"):
        section_overrides_with_suggestions(tailored, [1])


def test_exports_accept_webp_photo_and_strip_invalid_control_characters() -> None:
    image_stream = BytesIO()
    Image.new("RGB", (80, 100), "white").save(image_stream, format="WEBP")
    photo = "data:image/webp;base64," + base64.b64encode(image_stream.getvalue()).decode()
    resume = _resume().model_copy(update={"summary": "可靠\x0b负责的开发者"})

    word = render_resume_docx(resume, photo_asset=photo)
    pdf = render_resume_pdf(resume, photo_asset=photo)

    assert word.startswith(b"PK")
    assert pdf.startswith(b"%PDF-")
